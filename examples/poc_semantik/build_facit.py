#!/usr/bin/env python3
"""Semantik-PoC byggetrin 1: facit fra ekspertens fulde linje-gennemgang.

Læser arket "Alle linjer" i kundens Excel-gennemgang (28 kolonner, én linje
pr. bogført købsposteringslinje, MED ekspertens klassifikation) og skriver et
generisk facit-JSON: pr. linje id, konto+navn, leverandør, tekst, momskode,
beløb, moms og ekspert-fund-id (eller "OK").

Kundefilen ligger UDENFOR repoet og røres aldrig (§2a-princip: ren læsning).
Facit-JSON'en skrives til --out, som skal pege UDENFOR repoet (scratchpad) —
den indeholder posteringstekster, leverandørnavne og beløb og må ALDRIG
committes.

Brug:
  python3 build_facit.py --xlsx <sti/til/AI\\ Transaktionsgennemgang.xlsx> \\
      --out <scratchpad>/facit.json
"""
from __future__ import annotations

import argparse
import datetime
import json
from collections import Counter
from pathlib import Path

import openpyxl

SHEET_NAME = "Alle linjer"

# Forventede kolonneoverskrifter (position er ligegyldig — vi slår op på navn,
# så et evt. ombygget ark ikke stille tavst forskyder felter).
COLUMN_MAP = {
    "Bilagsnr.": "bilagsnr",
    "Posteringsdato": "posting_date",
    "Momsdato": "vat_date",
    "Konto": "account",
    "Kontonavn": "account_name",
    "Posteringstekst": "text",
    "Leverandør": "supplier",
    "Eksternt bilagsnr.": "external_doc_no",
    "Valuta": "currency",
    "Beløb ekskl. moms": "amount_excl_vat",
    "Momsbeløb": "vat_amount",
    "Ikke-fradragsber. moms": "non_deductible_vat",
    "Effektiv momssats": "effective_vat_rate",
    "Momsvirksomhedsgruppe": "vat_business_group",
    "Momsvaregruppe": "vat_product_group",
    "Risikoniveau": "risk_level",
    "Fund-ID": "fund_id",
    "Kategori": "category",
    "Lovgrundlag": "lovgrundlag",
    "Kommentar": "comment",
    "Anbefalet handling": "anbefalet_handling",
}


def _cell(v):
    """JSON-sikker konvertering (datetime -> ISO-dato)."""
    if isinstance(v, (datetime.datetime, datetime.date)):
        return v.isoformat()
    return v


def load_lines(xlsx_path: Path) -> list[dict]:
    wb = openpyxl.load_workbook(xlsx_path, read_only=True, data_only=True)
    if SHEET_NAME not in wb.sheetnames:
        raise SystemExit(f"Fandt ikke arket '{SHEET_NAME}' i {xlsx_path}. Arke: {wb.sheetnames}")
    ws = wb[SHEET_NAME]
    rows = ws.iter_rows(values_only=True)
    header = next(rows)
    idx_by_field = {}
    for i, name in enumerate(header):
        if name in COLUMN_MAP:
            idx_by_field[COLUMN_MAP[name]] = i
    missing = set(COLUMN_MAP.values()) - set(idx_by_field.keys())
    if missing:
        raise SystemExit(f"Mangler forventede kolonner i arket: {sorted(missing)}")

    lines = []
    for row_num, row in enumerate(rows, start=2):
        if row is None or all(v is None for v in row):
            continue
        rec = {field: _cell(row[i]) for field, i in idx_by_field.items()}
        fund_id = (rec.get("fund_id") or "").strip() if isinstance(rec.get("fund_id"), str) else rec.get("fund_id")
        rec["fund_id"] = fund_id if fund_id else "OK"
        rec["line_id"] = f"L{row_num:04d}"
        rec["vat_code"] = f"{rec.get('vat_business_group') or ''}/{rec.get('vat_product_group') or ''}"
        # Fradragsstatus: afledt signal LLM'en får at se (uden at afsløre facit).
        va = rec.get("vat_amount") or 0
        rec["moms_fratrukket"] = bool(va and va != 0)
        lines.append(rec)
    return lines


def build_meta(lines: list[dict], xlsx_path: Path) -> dict:
    fund_dist = Counter(l["fund_id"] for l in lines)
    risk_dist = Counter(l.get("risk_level") for l in lines)
    total_excl = round(sum(l.get("amount_excl_vat") or 0 for l in lines), 2)
    total_vat = round(sum(l.get("vat_amount") or 0 for l in lines), 2)
    total_non_ded = round(sum(l.get("non_deductible_vat") or 0 for l in lines), 2)
    return {
        "generated_at": datetime.datetime.now().isoformat(timespec="seconds"),
        "source_file": str(xlsx_path),
        "source_sheet": SHEET_NAME,
        "total_lines": len(lines),
        "fund_distribution": dict(sorted(fund_dist.items(), key=lambda kv: -kv[1])),
        "risk_distribution": dict(risk_dist),
        "kontrolsummer": {
            "beloeb_ekskl_moms": total_excl,
            "momsbeloeb": total_vat,
            "ikke_fradragsber_moms": total_non_ded,
            "note": (
                "Sammenlign disse tre summer mod kundefilens eget afstemningsgrundlag "
                "i arket 'Metode og forudsætninger' §1 (Datagrundlag og afstemning). "
                "Afvigelse > få øre indikerer kolonnefejl eller filtrerede rækker."
            ),
        },
    }


def main():
    ap = argparse.ArgumentParser(description=__doc__)
    ap.add_argument("--xlsx", required=True, type=Path, help="Sti til kunde-Excel (UDENFOR repo)")
    ap.add_argument("--out", required=True, type=Path, help="Facit-JSON output (UDENFOR repo, fx scratchpad)")
    args = ap.parse_args()

    lines = load_lines(args.xlsx)
    meta = build_meta(lines, args.xlsx)
    args.out.parent.mkdir(parents=True, exist_ok=True)
    args.out.write_text(json.dumps({"meta": meta, "lines": lines}, ensure_ascii=False, indent=2))

    print(f"Facit skrevet: {args.out}  ({meta['total_lines']} linjer)")
    print("Fund-fordeling (top 10):")
    for fid, n in list(meta["fund_distribution"].items())[:10]:
        print(f"  {fid}: {n}")
    print("Kontrolsummer:", meta["kontrolsummer"])


if __name__ == "__main__":
    main()
