"""
report_workbook.py — rådgiverens Excel-arbejdsbilag til kunderapporten
(byggetrin ~10, Bal-godkendt 2026-09-18, openpyxl — allerede i
``requirements.txt``).

Modsat den kundevendte HTML-rapport (``tools/generate_report.py``) er dette
arbejdsbilaget: HER (og i HTML-rapportens lineage-footer) må kontrolnumre
optræde. Filen indeholder ALLE fund (alle severities), ikke kun de
kuraterede/forfremmede — den er rådgiverens fulde evidensgrundlag, ikke et
kundedokument. Del/gem den derfor med samme forsigtighed som selve
rapport-JSON'en (kan indeholde kundedata).

Ark:
    1. "Oversigt"      — nøgletal + tema-tabel (tema, kildekontroller, antal,
                          beløb, medtag i kunderapporten).
    2. "Alle fund"      — flad liste, ALLE severities, kontrol-id, kontrolnavn,
                          tema (tom hvis kontrollen ikke er tema-mappet),
                          konto+navn, bilag, dato, beløb, beskrivelse.
    3. Ét ark pr. tema  — samme kolonner, filtreret til temaets kontroller.
"""

from __future__ import annotations

import re

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

from tools import report_themes

_HEADER_FILL = PatternFill(start_color="FF1B365D", end_color="FF1B365D", fill_type="solid")
_HEADER_FONT = Font(color="FFFFFFFF", bold=True)

_FINDING_COLUMNS = (
    "Kontrol-id", "Kontrolnavn", "Tema", "Severity", "Konto", "Kontonavn",
    "Bilag", "Dato", "Beløb", "Beskrivelse", "Antal referencer",
)


def _sheet_name(name: str, used: set) -> str:
    """Excel-ark-navne: max 31 tegn, ingen af : \\ / ? * [ ]. Sikrer unikke
    navne (tilføjer et løbenummer ved kollision efter afkortning)."""
    cleaned = re.sub(r"[:\\/?*\[\]]", "-", name).strip() or "Ark"
    base = cleaned[:31]
    candidate = base
    n = 2
    while candidate in used:
        suffix = f" ({n})"
        candidate = base[: 31 - len(suffix)] + suffix
        n += 1
    used.add(candidate)
    return candidate


def _finding_row(finding: dict, konto_navne: dict) -> tuple:
    refs = finding.get("transactions") or []
    first = refs[0] if refs else {}
    konto = first.get("account_id", "") or ""
    return (
        finding.get("test_id"),
        finding.get("test_name", ""),
        report_themes.THEMES.get(report_themes.theme_of(finding.get("test_id")) or "", {}).get("navn", ""),
        finding.get("severity", ""),
        konto,
        konto_navne.get(konto, ""),
        first.get("transaction_id") or first.get("journal_id") or "",
        first.get("date", "") or "",
        finding.get("estimated_amount") or 0.0,
        finding.get("description", ""),
        len(refs),
    )


def _write_findings_sheet(ws, findings: list, konto_navne: dict) -> None:
    ws.append(_FINDING_COLUMNS)
    for col in range(1, len(_FINDING_COLUMNS) + 1):
        cell = ws.cell(row=1, column=col)
        cell.font = _HEADER_FONT
        cell.fill = _HEADER_FILL
    ws.freeze_panes = "A2"
    for f in findings:
        ws.append(_finding_row(f, konto_navne))
    widths = (10, 42, 24, 10, 12, 28, 16, 12, 14, 60, 10)
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w


def _write_overview_sheet(ws, report: dict, analytics: dict, curation: dict | None) -> None:
    ws.append(["VAT Analytics — arbejdsbilag (rådgiver)"])
    ws["A1"].font = Font(bold=True, size=14)
    lineage = report.get("lineage") or {}
    ws.append(["Katalogversion", lineage.get("catalog_version", "")])
    ws.append(["Datakontrakt-version", lineage.get("data_contract_version", "")])
    ws.append(["Kørselstidspunkt (UTC)", lineage.get("generated_at", "")])
    ws.append([])

    sev = analytics.get("severity_summary", {}) or {}
    ws.append(["Fund i alt", analytics.get("total_findings", 0)])
    ws.append(["Kritisk", sev.get("critical", 0)])
    ws.append(["Høj", sev.get("high", 0)])
    ws.append(["Medium", sev.get("medium", 0)])
    ws.append(["Lav", sev.get("low", 0)])
    ws.append([])

    ws.append(["Tema", "Kildekontroller", "Antal fund", "Beløb i alt", "Medtaget i kunderapport", "Status"])
    header_row = ws.max_row
    for col in range(1, 7):
        cell = ws.cell(row=header_row, column=col)
        cell.font = _HEADER_FONT
        cell.fill = _HEADER_FILL

    groups = (curation or {}).get("grupper", {})
    for tk in report_themes.THEME_ORDER:
        grp = groups.get(tk)
        meta = report_themes.THEMES[tk]
        auto = (grp or {}).get("auto", {})
        ws.append([
            meta["navn"],
            ", ".join(str(c) for c in meta["kontroller"]),
            auto.get("fund_antal", 0),
            auto.get("beloeb_i_alt", 0.0),
            "Ja" if (grp or {}).get("medtag") else "Nej",
            (grp or {}).get("status", "(ingen fund)"),
        ])
    for i, w in enumerate((28, 20, 12, 16, 22, 26), start=1):
        ws.column_dimensions[get_column_letter(i)].width = w


def build_workbook(report: dict, curation: dict | None, out_path: str) -> None:
    """Byg og gem arbejdsbilaget. Fejler ALDRIG stille på tomt datagrundlag
    (samme disciplin som resten af rapport-laget) — et tomt fundgrundlag
    giver blot tomme ark, ikke en exception."""
    analytics = report.get("analytics") if isinstance(report.get("analytics"), dict) else report
    all_findings = analytics.get("all_findings") or []
    konto_navne = report.get("konto_navne") or {}

    wb = Workbook()
    overview_ws = wb.active
    overview_ws.title = "Oversigt"
    _write_overview_sheet(overview_ws, report, analytics, curation)

    all_ws = wb.create_sheet("Alle fund")
    _write_findings_sheet(all_ws, all_findings, konto_navne)

    used_names = {"Oversigt", "Alle fund"}
    for tk in report_themes.THEME_ORDER:
        controls = report_themes.THEMES[tk]["kontroller"]
        theme_findings = [f for f in all_findings if f.get("test_id") in controls]
        if not theme_findings:
            continue
        sheet_name = _sheet_name(f"Tema {report_themes.THEMES[tk]['navn']}", used_names)
        ws = wb.create_sheet(sheet_name)
        _write_findings_sheet(ws, theme_findings, konto_navne)

    wb.save(out_path)
