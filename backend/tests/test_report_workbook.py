"""
tools/report_workbook.py — rådgiverens Excel-arbejdsbilag (byggetrin ~10):
ALLE fund (alle severities), kontrol-id'er tilladt (modsat kundevisningen).
"""

from openpyxl import load_workbook

from tools import report_curation as rc
from tools import report_workbook as rw


def _finding(test_id, severity="medium", account_id="4000", amount=100.0, test_name="Testkontrol"):
    return {
        "test_id": test_id, "test_name": test_name, "severity": severity,
        "estimated_amount": amount, "description": "beskrivelse",
        "transactions": [{"account_id": account_id, "transaction_id": "T1",
                           "date": "2025-01-10", "amount": amount}],
    }


def _report(findings, konto_navne=None):
    return {
        "lineage": {"catalog_version": "1.3.0", "data_contract_version": "0.4.4",
                    "generated_at": "2026-09-18T10:00:00+00:00"},
        "konto_navne": konto_navne or {},
        "analytics": {
            "all_findings": findings,
            "total_findings": len(findings),
            "severity_summary": {
                "critical": sum(1 for f in findings if f["severity"] == "critical"),
                "high": sum(1 for f in findings if f["severity"] == "high"),
                "medium": sum(1 for f in findings if f["severity"] == "medium"),
                "low": sum(1 for f in findings if f["severity"] == "low"),
            },
        },
    }


def test_build_workbook_creates_overview_and_all_findings_sheets(tmp_path):
    findings = [_finding(19, severity="high"), _finding(63, severity="low")]  # 63 = ikke tema-mappet
    report = _report(findings, konto_navne={"4000": "Salgskonto"})
    curation = rc.seed_curation(findings)
    out_path = tmp_path / "arbejdsbilag.xlsx"

    rw.build_workbook(report, curation, str(out_path))
    assert out_path.exists()

    wb = load_workbook(str(out_path))
    assert "Oversigt" in wb.sheetnames
    assert "Alle fund" in wb.sheetnames

    all_ws = wb["Alle fund"]
    assert all_ws.cell(row=1, column=1).value == "Kontrol-id"
    # 2 fund + header = 3 rækker
    assert all_ws.max_row == 3
    values = {all_ws.cell(row=r, column=1).value for r in (2, 3)}
    assert values == {19, 63}


def test_build_workbook_includes_konto_navn_lookup(tmp_path):
    findings = [_finding(19, account_id="4000")]
    report = _report(findings, konto_navne={"4000": "Salgskonto"})
    out_path = tmp_path / "arbejdsbilag.xlsx"
    rw.build_workbook(report, None, str(out_path))

    wb = load_workbook(str(out_path))
    all_ws = wb["Alle fund"]
    header = [all_ws.cell(row=1, column=c).value for c in range(1, all_ws.max_column + 1)]
    konto_navn_col = header.index("Kontonavn") + 1
    assert all_ws.cell(row=2, column=konto_navn_col).value == "Salgskonto"


def test_build_workbook_creates_one_sheet_per_theme_with_findings(tmp_path):
    findings = [_finding(19, severity="high"), _finding(80, severity="medium")]
    report = _report(findings)
    curation = rc.seed_curation(findings)
    out_path = tmp_path / "arbejdsbilag.xlsx"
    rw.build_workbook(report, curation, str(out_path))

    wb = load_workbook(str(out_path))
    theme_sheets = [n for n in wb.sheetnames if n.startswith("Tema ")]
    assert any("Kodeopsætning" in n for n in theme_sheets)
    assert any("Momsbehandling" in n for n in theme_sheets)
    # Intet tema-ark for temaer uden fund i denne kørsel:
    assert not any("Timing" in n for n in theme_sheets)


def test_build_workbook_empty_findings_does_not_crash(tmp_path):
    report = _report([])
    out_path = tmp_path / "arbejdsbilag.xlsx"
    rw.build_workbook(report, None, str(out_path))
    assert out_path.exists()
    wb = load_workbook(str(out_path))
    assert wb["Alle fund"].max_row == 1  # kun header


def test_sheet_name_truncates_and_deduplicates():
    used = set()
    name1 = rw._sheet_name("Tema " + "x" * 40, used)
    name2 = rw._sheet_name("Tema " + "x" * 40, used)
    assert len(name1) <= 31
    assert len(name2) <= 31
    assert name1 != name2


def test_sheet_name_strips_invalid_excel_characters():
    used = set()
    name = rw._sheet_name("Tema: a/b?c*d[e]", used)
    assert not any(ch in name for ch in ":\\/?*[]")
