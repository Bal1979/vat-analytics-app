"""
Multi-entity-dimensionen (vat-extract Del 10, grænsefladen til vat-analytics;
Bal-godkendt 2026-10-02, datakontrakt v0.6.0).

Dækker:
  1. Kontrakten: ``lines.entity_id`` er en balai_extension (v0.6.0).
  2. Parseren læser ``ext_entity_id`` -> ``lines[].entity_id`` (nøglesæt-
     symmetri: altid til stede, "" ved fravær) -- i dag ignoreredes kolonnen.
  3. Det ÆRLIGE VÆRN: >= 2 distinkte entity_id uden præfikset bilagsnøgle
     (ingen '|' i invoice_numbers) giver en tydelig advarsel i parse_info
     (+ rapporten) i stedet for stille at regne videre. Ingen automatisk
     omnøgling.
  4. Downstream: transaktions-id'er med '|' (``DOC_VAND|1001_<dato>``)
     knækker hverken motor, curation, HTML-rapport, Excel-arbejdsbilag eller
     JSON-eksport.

Alle fixtures er opdigtede (enheds-id'er VAND/ENER er syntetiske) -- ingen
kundedata.
"""

import csv
import json

from openpyxl import load_workbook

from parsers import canonical_parser
from parsers import data_adapter
from tools import analyze_canonical
from tools import build_data_contract as gen
from tools import data_contract_data as dcd
from tools import generate_report, report_curation, report_workbook


_BASE = [
    "posting_dates", "tax_point", "vat_period", "credit_note_flag",
    "invoice_numbers", "gl_accounts", "vat_codes", "vat_amount",
    "debit_amount", "credit_amount", "supply_direction", "currency_fx",
]


def _row(date, inv, acct="6000", code="DOMESTIC|U25", vat="250.0", debit="1000.0",
         credit="0", entity=None):
    r = [date, date, date[:7], "false", inv, acct, code, vat, debit, credit, "purchase", ""]
    if entity is not None:
        r.append(entity)
    return r


def _write(tmp_path, rows, with_entity=True, name="gl_entries.csv"):
    path = tmp_path / name
    headers = _BASE + (["ext_entity_id"] if with_entity else [])
    with open(path, "w", encoding="utf-8", newline="") as f:
        w = csv.writer(f)
        w.writerow(headers)
        w.writerows(rows)
    return str(path)


def _parse(tmp_path, rows, with_entity=True):
    canonical, info = canonical_parser.parse_canonical(_write(tmp_path, rows, with_entity))
    assert canonical is not None, info
    return canonical, info


def _lines(canonical):
    return [ln for t in canonical["transactions"] for ln in t["lines"]]


def _warn(info):
    return [w for w in info["warnings"] if "Multi-entity" in w]


# --- 1. Kontrakt -----------------------------------------------------------------

def test_contract_has_entity_id_as_balai_extension_v060():
    # entity_id blev indført i v0.6.0; kontrakten er siden bumpet additivt (v0.7.0).
    assert tuple(int(x) for x in dcd.CONTRACT_VERSION.split(".")) >= (0, 6, 0)
    contract, _problems = gen.build_contract()
    lines = contract["objekter"]["transactions"]["sub_objekt"]["felter"]
    field = next(f for f in lines if f["navn"] == "entity_id")
    assert field["ekstension"] is True
    assert field["obligatorisk"] is False
    assert field["kilder"] == {"excel": False, "saft": False, "canonical": "partial"}
    assert "AGGREGERING" in field["kraeves_af"] and "SEGMENTERING" in field["kraeves_af"]
    ext = [e for e in contract["balai_extensions"]["felter"] if e["felt"] == "entity_id"]
    assert len(ext) == 1 and "SAF-T Financial" in ext[0]["begrundelse"]


# --- 2. Parseren -----------------------------------------------------------------

def test_parser_reads_ext_entity_id_onto_lines(tmp_path):
    canonical, info = _parse(tmp_path, [
        _row("2025-01-10", "VAND|1001", entity="VAND"),
        _row("2025-01-11", "ENER|1001", entity="ENER"),
    ])
    ids = sorted(ln["entity_id"] for ln in _lines(canonical))
    assert ids == ["ENER", "VAND"]
    assert info["multi_entity"]["entity_ids"] == ["ENER", "VAND"]
    assert info["multi_entity"]["antal_enheder"] == 2


def test_parser_entity_id_is_trimmed(tmp_path):
    canonical, _ = _parse(tmp_path, [_row("2025-01-10", "VAND|1", entity="  VAND ")])
    assert _lines(canonical)[0]["entity_id"] == "VAND"


def test_parser_without_column_gives_empty_entity_id_and_neutral_info(tmp_path):
    canonical, info = _parse(tmp_path, [_row("2025-01-10", "F1"), _row("2025-01-11", "F2")],
                              with_entity=False)
    assert all(ln["entity_id"] == "" for ln in _lines(canonical))
    me = info["multi_entity"]
    assert me["antal_enheder"] == 0 and me["kollisionsrisiko"] is False
    assert _warn(info) == []


def test_parser_blank_entity_cells_are_empty_string(tmp_path):
    canonical, info = _parse(tmp_path, [_row("2025-01-10", "F1", entity="")])
    assert _lines(canonical)[0]["entity_id"] == ""
    assert info["multi_entity"]["antal_enheder"] == 0


def test_other_input_paths_keep_key_set_symmetry():
    # Excel-vejen (data_adapter) og SAF-T-vejen bærer nøglen som "".
    parsed = {
        "header": {}, "accounts": [], "tax_table": [], "suppliers": [], "customers": [],
        "parse_info": {},
        "transactions": [{
            "transaction_id": "T1", "date": "2025-01-10", "account_id": "6000",
            "debit_amount": 100.0, "credit_amount": 0.0, "vat_amount": 25.0,
            "vat_code": "U25", "invoice_number": "F1",
        }],
    }
    out = data_adapter.adapt_excel_to_saft(parsed)
    assert out["transactions"][0]["lines"][0]["entity_id"] == ""


# --- 3. Det ærlige værn ----------------------------------------------------------

def test_guard_warns_for_two_entities_without_prefixed_keys(tmp_path):
    # Samme (bilagsnr, dato) i to enheder -> falsk sammensmeltning.
    canonical, info = _parse(tmp_path, [
        _row("2025-01-10", "1001", entity="VAND"),
        _row("2025-01-10", "1001", entity="ENER"),
        _row("2025-01-12", "1002", entity="VAND"),
    ])
    me = info["multi_entity"]
    assert me["kollisionsrisiko"] is True
    assert me["linjer_uden_praefiks"] == 3
    assert me["bilagsnoegler_i_flere_enheder"] == 1
    warns = _warn(info)
    assert len(warns) == 1
    assert "Multi-entity uden præfikset bilagsnøgle" in warns[0]
    assert "kollisionsrisiko" in warns[0]
    assert "--praefiks-bilagsnoegle" in warns[0]
    # Motoren omnøgler IKKE: de to enheders bilag er (som advaret) smeltet sammen.
    assert canonical["summary"]["total_transactions"] == 2


def test_guard_warns_even_when_no_concrete_collision_measured(tmp_path):
    _, info = _parse(tmp_path, [
        _row("2025-01-10", "1001", entity="VAND"),
        _row("2025-01-11", "2002", entity="ENER"),
    ])
    assert info["multi_entity"]["kollisionsrisiko"] is True
    assert info["multi_entity"]["bilagsnoegler_i_flere_enheder"] == 0
    assert "Ingen konkret kollision er målt" in _warn(info)[0]


def test_guard_silent_when_keys_are_prefixed(tmp_path):
    canonical, info = _parse(tmp_path, [
        _row("2025-01-10", "VAND|1001", entity="VAND"),
        _row("2025-01-10", "ENER|1001", entity="ENER"),
    ])
    assert info["multi_entity"]["kollisionsrisiko"] is False
    assert info["multi_entity"]["bilagsnoegler_i_flere_enheder"] == 0
    assert _warn(info) == []
    # Præfikset holder de to enheders identiske bilagsnr adskilt.
    assert canonical["summary"]["total_transactions"] == 2
    tids = sorted(t["transaction_id"] for t in canonical["transactions"])
    assert tids == ["DOC_ENER|1001_2025-01-10", "DOC_VAND|1001_2025-01-10"]


def test_guard_warns_for_partially_prefixed_keys(tmp_path):
    _, info = _parse(tmp_path, [
        _row("2025-01-10", "VAND|1001", entity="VAND"),
        _row("2025-01-10", "1001", entity="ENER"),
    ])
    me = info["multi_entity"]
    assert me["kollisionsrisiko"] is True and me["linjer_uden_praefiks"] == 1
    assert "delvist" in _warn(info)[0]


def test_guard_silent_for_single_entity_even_unprefixed(tmp_path):
    _, info = _parse(tmp_path, [
        _row("2025-01-10", "1001", entity="VAND"),
        _row("2025-01-11", "1002", entity="VAND"),
    ])
    assert info["multi_entity"]["antal_enheder"] == 1
    assert info["multi_entity"]["kollisionsrisiko"] is False
    assert _warn(info) == []


def test_guard_ignores_lines_without_invoice_number(tmp_path):
    # Linjer uden bilagsnr nøgles ikke (hver sin gruppe) -> ingen kollisionsrisiko.
    _, info = _parse(tmp_path, [
        _row("2025-01-10", "", entity="VAND"),
        _row("2025-01-10", "", entity="ENER"),
    ])
    assert info["multi_entity"]["antal_enheder"] == 2
    assert info["multi_entity"]["kollisionsrisiko"] is False
    assert _warn(info) == []


def test_guard_counts_lines_without_entity_when_entities_present(tmp_path):
    _, info = _parse(tmp_path, [
        _row("2025-01-10", "VAND|1", entity="VAND"),
        _row("2025-01-10", "ENER|1", entity="ENER"),
        _row("2025-01-10", "X1", entity=""),
    ])
    assert info["multi_entity"]["linjer_uden_entity_id"] == 1


# --- 4. Rapport-/CLI-laget -------------------------------------------------------

def test_report_carries_warning_and_html_shows_forbehold(tmp_path):
    path = _write(tmp_path, [
        _row("2025-01-10", "1001", entity="VAND"),
        _row("2025-01-10", "1001", entity="ENER"),
    ])
    report = analyze_canonical.build_report(path, None, None, 0.01, "alle")
    pi = report["parse_info"]
    assert pi["multi_entity"]["kollisionsrisiko"] is True
    assert any("Multi-entity uden præfikset bilagsnøgle" in w for w in pi["warnings"])
    html = generate_report.render_html(report, None, niveau=3)
    assert "flere regnskabsenheder uden præfikset bilagsnøgle" in html.lower()


def test_report_html_has_no_forbehold_when_prefixed_or_single_entity(tmp_path):
    path = _write(tmp_path, [
        _row("2025-01-10", "VAND|1001", entity="VAND"),
        _row("2025-01-10", "ENER|1001", entity="ENER"),
    ])
    report = analyze_canonical.build_report(path, None, None, 0.01, "alle")
    assert report["parse_info"]["multi_entity"]["kollisionsrisiko"] is False
    html = generate_report.render_html(report, None, niveau=3)
    assert "uden præfikset bilagsnøgle" not in html.lower()


def test_cli_summary_mentions_multi_entity(tmp_path, capsys):
    path = _write(tmp_path, [
        _row("2025-01-10", "1001", entity="VAND"),
        _row("2025-01-10", "1001", entity="ENER"),
    ])
    report = analyze_canonical.build_report(path, None, None, 0.01, "alle")
    analyze_canonical._print_summary(report)
    out = capsys.readouterr().out
    assert "Multi-entity: 2 regnskabsenhed(er)" in out and "KOLLISIONSRISIKO" in out


# --- 5. Downstream: '|' i transaktions-id'er -------------------------------------

def _prefixed_dataset(tmp_path):
    rows = []
    for ent in ("VAND", "ENER"):
        for i in range(1, 25):
            d = f"2025-01-{i % 28 + 1:02d}"
            rows.append(_row(d, f"{ent}|{1000 + i}", vat="250.0" if i % 3 else "1750.0", entity=ent))
            rows.append(_row(d, f"{ent}|{1000 + i}", acct="1900", code="", vat="0",
                             debit="0", credit="1250.0", entity=ent))
    # Dublet-bilagsnr i samme enhed -> giver fund med '|' i transaktions-id'et.
    rows.append(_row("2025-02-01", "VAND|1001", entity="VAND"))
    rows.append(_row("2025-02-05", "VAND|1001", entity="VAND"))
    return _write(tmp_path, rows)


def test_pipe_in_transaction_ids_does_not_break_report_workbook_or_export(tmp_path):
    path = _prefixed_dataset(tmp_path)
    report = analyze_canonical.build_report(path, None, None, 0.01, "alle")
    assert "fejl" not in report
    findings = report["analytics"]["all_findings"]
    pipe_ids = {
        t["transaction_id"] for f in findings for t in f.get("transactions", [])
        if "|" in str(t.get("transaction_id", ""))
    }
    assert pipe_ids, "fixturen skal give fund der refererer id'er med '|'"
    assert all(i.startswith("DOC_") for i in pipe_ids)

    # JSON-eksport (rapport-JSON'en) -- rundtur uden tab.
    dumped = json.dumps(report, ensure_ascii=False)
    back = json.loads(dumped)
    back_ids = {
        t["transaction_id"] for f in back["analytics"]["all_findings"]
        for t in f.get("transactions", []) if "|" in str(t.get("transaction_id", ""))
    }
    assert back_ids == pipe_ids

    # Curation + HTML-rapport (niveau 3 + appendix) + Excel-arbejdsbilag.
    curation = report_curation.seed_curation(findings)
    html = generate_report.render_html(report, None, niveau=3, curation=curation, appendix=True)
    assert "<html" in html
    out = tmp_path / "arbejdsbilag.xlsx"
    report_workbook.build_workbook(report, curation, str(out))
    wb = load_workbook(str(out))
    assert "Alle fund" in wb.sheetnames


def test_pipe_ids_survive_curation_html_and_workbook_with_synthetic_findings(tmp_path):
    tid = "DOC_VAND|1001_2025-01-10"
    findings = [{
        "test_id": 19, "test_name": "Testkontrol", "severity": "high",
        "estimated_amount": 100.0, "description": "beskrivelse",
        "transactions": [{"account_id": "6000", "transaction_id": tid,
                           "date": "2025-01-10", "amount": 100.0}],
    }]
    report = {
        "lineage": {"catalog_version": "1.5.1", "data_contract_version": "0.6.0",
                    "generated_at": "2026-10-02T10:00:00+00:00"},
        "konto_navne": {}, "parse_info": {"sections": {}, "multi_entity": {}},
        "analytics": {
            "all_findings": findings, "total_findings": 1,
            "severity_summary": {"critical": 0, "high": 1, "medium": 0, "low": 0},
            "summary": {"currency": "DKK"},
        },
    }
    curation = report_curation.seed_curation(findings)
    html = generate_report.render_html(report, None, niveau=3, curation=curation, appendix=True)
    assert "<html" in html
    out = tmp_path / "wb.xlsx"
    report_workbook.build_workbook(report, curation, str(out))
    ws = load_workbook(str(out))["Alle fund"]
    assert ws.max_row == 2
    assert json.loads(json.dumps(findings))[0]["transactions"][0]["transaction_id"] == tid
